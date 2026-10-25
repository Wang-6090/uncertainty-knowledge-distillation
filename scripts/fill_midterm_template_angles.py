"""Generate a midterm report whose sections follow the template's required angles."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "附件2-北京交通大学大学生创新创业训练计划项目中期进展报告模板.docx"
OUTPUT = ROOT / "附件2-北京交通大学大学生创新创业训练计划项目中期进展报告_按模板内容角度版.docx"
PROJECT_NAME = "基于不确定性知识蒸馏的新类发现方法"


# The template requests: background, problem, technology, solved problem, expected effect.
INTRO = [
    "研究背景是：智能视觉系统面向开放环境，测试输入可能包含训练阶段未纳入的类别。封闭集分类模型会将其误判为已知类别，难以完成未知识别和新类发现。本项目拟解决在缺少未知类别标注时，联合完成已知类别分类、未知类别拒识和新类别聚类的问题。",
    "项目采用教师—学生知识蒸馏、教师不确定性建模、独立拒识表征、非负正—未标记学习和无监督聚类等技术：利用教师软类别分布指导学生学习，利用不确定性调节蒸馏强度，利用拒识分支从混合无标签数据中学习未知风险，并对候选样本进行聚类。",
    "已建立 CIFAR-100 语义困难 60/40 实验协议并完成多组对照。结果表明，shared-backbone 拒识表征能够改善未知类别拒识，不确定性融合在部分划分上有效，但不确定性直接蒸馏尚未形成稳定增益。预期形成可复现实验程序、研究报告，并进一步改善语义相近类别的表征重叠。",
]


# The template requests a list of key problems or the research scheme, not a paper-style review.
RESEARCH_BLOCKS = [
    (
        "（一）主要研究内容",
        [
            "1. 重点研究问题。第一，教师模型的不确定性是否能够识别不可靠知识，并通过蒸馏改善学生模型对未知类别的区分能力；第二，在混合发现池没有未知类别标签的条件下，如何学习稳定的未知类别拒识表征；第三，如何提高困难未知类别进入候选池的覆盖率，并将候选样本正确聚类为新类别。",
            "2. 研究方案。建立“教师模型—学生模型—拒识分支—候选聚类”的处理流程。教师和学生模型负责已知类别分类与知识蒸馏；独立 rejection encoder 负责提取拒识特征；NNPU 目标利用已知标注样本和混合未标记样本估计拒识风险；已知验证集用于模型选择和拒识阈值校准；KMeans 用于候选未知样本聚类，auto K 用于估计新类别数量。",
            "3. 不确定性知识蒸馏。以标准 logits-KD 为对照，使用温度缩放后的 KL 散度传递教师模型的软类别分布，并保留已知类别交叉熵。将不确定性 u_i 转换为 exp(-u_i) 样本权重，并进行 batch-mean normalization。该修改用于降低教师模型对边界样本或不可靠预测的过度传递；同时比较 learned uncertainty 和 MC-epistemic 两种不确定性来源。",
            "4. 拒识表征与 NNPU。除学生分类分支外增加独立 rejection encoder，使模型分别学习类别判别特征和未知风险特征。拒识训练采用 NNPU，不将混合发现池的真实类别作为训练标签；拒识编码器采用 shared-backbone 初始化，即复制学生 ResNet-18 的表征参数后独立优化。该修改针对 scratch CNN 拒识分支表征能力不足的问题。",
            "5. 候选筛选与聚类。依据已知验证集的目标已知类别覆盖率确定拒识阈值，将高风险样本形成候选未知样本池；采用 PCA 降维和 KMeans 聚类；比较 projection PCA 与 feature PCA 对聚类表示的影响，并比较 prototype-boundary 候选加权对候选筛选的影响。",
            "6. 实验对照与评价。标准 KD 与 uncertainty-KD 用于评价训练期不确定性蒸馏；scratch、random-backbone 与 shared-backbone 用于评价拒识表征初始化；feature-rejector 与 validation-only uncertainty fusion 用于评价不确定性后处理。所有实验同时报告 AUROC、AUPR、FPR95、OSCR、已知类别接受率、未知类别拒识率、接受后已知类别准确率、分数重叠、自动 K 和 ACC/NMI/ARI。",
            "7. 与立项内容的调整。立项时拟将 uncertainty-KD 作为主要训练机制。阶段性实验显示，uncertainty-KD 在不同数据划分上的结果不稳定，未形成完整新类发现增益。因此将研究重点调整为“标准 KD 对照下的拒识表征与不确定性拒识融合”，保留 uncertainty-KD 作为对照实验；后续进一步研究困难未知类别覆盖和类条件支持表征。",
        ],
    ),
]


PLAN = [
    "第 1 阶段（项目开始后第 1—2 个月，已完成）：开展相关文献和方法调研，明确研究问题与技术路线；完成 CIFAR-100 语义困难 60/40 数据划分、教师—学生训练程序和基础评价程序。",
    "第 2 阶段（第 3—5 个月，已完成）：完成交叉熵与标准知识蒸馏基线；实现不确定性知识蒸馏，比较 learned uncertainty 和 MC-epistemic 两种不确定性来源，并记录训练损失和权重变化。",
    "第 3 阶段（第 6—8 个月，已完成）：实现独立拒识编码器和 NNPU 拒识目标；完成 scratch、random-backbone 和 shared-backbone 对照；建立已知验证集模型选择与拒识阈值校准流程。",
    "第 4 阶段（第 9—10 个月，已完成主要实验）：完成 validation-only 不确定性融合、feature PCA、自动 K 和 prototype-boundary 候选加权实验；根据联合指标评价各方法是否具有完整任务收益。",
    "中期检查后：以标准 KD 和 shared-backbone 拒识表征为主线，研究仅使用已知训练集、已知验证集和模型预测信息的困难未知类别覆盖与类条件支持表征；完成最终多划分验证、实验总结和研究报告。",
]


PROGRESS_BLOCKS = [
    (
        "（一）已开展的工作",
        [
            "1. 完成文献与方法调研，围绕开放集识别、新类发现、知识蒸馏、不确定性估计、正—未标记学习和无监督聚类梳理技术路线，明确以“未知类别拒识和新类别聚类”为主要评价目标。",
            "2. 完成实验设计与基线构建。建立 CIFAR-100 语义困难 60/40 划分、5,400 个样本的混合发现池、教师—学生训练、已知验证集模型选择和拒识阈值校准流程；以标准 logits-KD 作为基本对照。",
            "3. 完成不确定性知识蒸馏实现与对照。增加 exp(-u) 样本权重和 batch-mean normalization，分别实现 learned uncertainty 与 MC-epistemic 路径，并检查权重、损失和梯度确实生效。2114—2116 三个新划分中，MC-epistemic 方法的未知类别拒识率平均提升 3.325 个百分点，但整体 NMI 平均下降 0.008623，未形成稳定增益。",
            "4. 完成独立拒识表征和 NNPU 实现。比较 scratch CNN、random-backbone 和 shared-backbone；scratch 分支在已知类别训练—验证探针任务上的准确率约为 18%—20%，shared-backbone 约为 73%—74%。2081—2083 三个新划分中，shared-backbone 相对 scratch 的平均 AUROC 提升 0.0320，OSCR 提升 0.0280，未知类别拒识率提升 3.61 个百分点。",
            "5. 完成不确定性拒识融合、候选筛选和聚类实验。7 个新划分中，validation-only uncertainty fusion 平均使 AUROC 提升 0.011677、未知类别拒识率提升 4.393 个百分点，但存在划分敏感性；feature PCA 提升 NMI/ARI；prototype-boundary 在 2119 中使未知类别拒识率下降 0.275 个百分点，未通过联合标准。",
            "6. 完成项目阶段性归档和程序检查。当前单元测试为 317 passed，Python 文件 compileall 通过；主要实验已保存数据划分、配置文件、模型 checkpoint、训练日志、校准报告、discovery 报告和结果分析文件。",
        ],
    ),
    (
        "（二）当前困难及解决办法",
        [
            "困难一：语义相近的未知类别与已知类别在特征空间和拒识分数上存在明显重叠。当前仍有约 74%—77% 的未知样本分数低于已知类别 q95，说明仅调整拒识阈值不能从根本上消除重叠。后续通过改进训练阶段的特征空间和类条件支持区域解决。",
            "困难二：不确定性可以排序部分已知分类错误，但尚未稳定区分未知类别。MC-epistemic logits-KD 在 2114—2116 三个新划分上使未知类别拒识率平均提高 3.325 个百分点，但整体 NMI 平均下降 0.008623，已知类别接受率平均下降 0.872 个百分点。后续停止不确定性权重函数搜索，将其保留为 validation-only 拒识辅助信号。",
            "困难三：候选池规模和候选规则的改变不必然提高困难未知类别覆盖率。prototype-boundary 在 2119 中改善排序指标，但未知类别拒识率下降 0.275 个百分点。后续同时评价候选污染率、未知覆盖率和逐类别覆盖率，避免仅依据排序指标调整候选规则。",
            "困难四：自动类别数估计和聚类表示具有划分敏感性。feature PCA 能够改善 NMI 和 ARI，但 Auto-K 误差在不同划分上方向不一致。后续保留 feature PCA 作为聚类模块，并将 Auto-K 误差作为独立问题研究，不以真实类别数量替代自动估计结果。",
        ],
    ),
]


STAGE_RESULTS = [
    "完成可复现实验流程。shared-backbone 在3个新划分中使AUROC平均提升0.0320、未知类别拒识率提升3.61个百分点；不确定性知识蒸馏未形成稳定增益。"
]


def set_font(paragraph, size=10.5, bold=False):
    for run in paragraph.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(size)
        run.bold = bold


def clear_cell(cell):
    cell.text = ""
    return cell.paragraphs[0]


def fill_cell(cell, blocks, size=10.5):
    clear_cell(cell)
    first = True
    for heading, paragraphs in blocks:
        if not first:
            cell.add_paragraph()
        p = cell.paragraphs[0] if first else cell.paragraphs[-1]
        p.text = heading
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p, size=size, bold=True)
        first = False
        for text in paragraphs:
            p = cell.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(p, size=size)


def fill_simple_cell(cell, paragraphs, size=10.5):
    clear_cell(cell)
    for i, text in enumerate(paragraphs):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph(text)
        if i == 0:
            p.text = text
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p, size=size)


def main():
    doc = Document(TEMPLATE)
    doc.paragraphs[6].text = f"项目名称：{PROJECT_NAME}"
    set_font(doc.paragraphs[6], size=18, bold=True)
    doc.tables[0].cell(0, 2).text = PROJECT_NAME
    set_font(doc.tables[0].cell(0, 2).paragraphs[0], size=10.5)
    doc.tables[0].cell(1, 2).text = "□研究报告  □软件程序"
    set_font(doc.tables[0].cell(1, 2).paragraphs[0], size=10.5)

    fill_simple_cell(doc.tables[1].cell(0, 0), INTRO, size=10.5)
    fill_cell(doc.tables[2].cell(0, 0), RESEARCH_BLOCKS + [("（二）实施进度安排", PLAN)], size=10.2)
    fill_cell(doc.tables[3].cell(0, 0), PROGRESS_BLOCKS, size=10.0)
    fill_simple_cell(doc.tables[4].cell(0, 0), STAGE_RESULTS, size=10.5)

    doc.core_properties.title = PROJECT_NAME + "中期进展报告"
    doc.core_properties.subject = "严格按模板内容角度填写"
    doc.core_properties.comments = "项目简介、研究内容、实施计划和实施进展均按模板规定角度填写；学生、成员和指导教师栏目保持模板原状。"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
