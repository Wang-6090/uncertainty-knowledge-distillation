"""Fill the experimental sections of the university midterm-report template."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "附件2-北京交通大学大学生创新创业训练计划项目中期进展报告模板.docx"
OUTPUT = ROOT / "附件2-北京交通大学大学生创新创业训练计划项目中期进展报告_规范表述版.docx"


PROJECT_NAME = "基于不确定性知识蒸馏的新类发现方法"

INTRO = [
    (
        "本项目研究新类发现问题：训练阶段仅有部分类别标注，测试阶段还会出现训练未见的新类别。"
        "研究目标是判断样本属于已知类或新类别，并对新类别样本进行聚类；考察教师模型不确定性"
        "能否提升学生模型的新类发现能力。"
    ),
    (
        "实验采用 CIFAR-100 semantic-hard 60/40 协议，将 60 个类别设为已知类、40 个类别设为新类。训练使用"
        "已知类标注样本和 5,400 个无标签混合发现池；系统学习已知类表征后，由独立 rejection 分支计算"
        "未知风险，并对被拒识样本进行聚类。阈值由已知验证集校准，使约 95% 的已知样本得到保留。"
    ),
    (
        "阶段性结果表明，不确定性能够部分反映已知分类错误，但直接用于知识蒸馏尚未形成稳定增益。"
        "当前主要困难是语义相近的新类与已知类的表征和分数重叠，下一阶段研究 hard-unknown 覆盖。"
    ),
]


RESEARCH_CONTENT = [
    (
        "（一）主要研究内容",
        [
            "1. 研究对象和数据：本项目研究开放集新类发现。已知类是训练阶段具有标注的类别，novel/unknown 是训练阶段没有标注、测试阶段出现的新类别，mixed discovery pool 是训练阶段仅提供图像、不提供类别标签的混合数据。实验采用 CIFAR-100 semantic-hard 60/40 划分，即 60 个已知类和 40 个新类；混合发现池固定为 5,400 个样本，已知先验为 0.2。",
            "2. 教师—学生知识蒸馏：教师模型作为参考模型，学生模型承担最终的新类发现任务。标准 logits-KD 使用温度缩放后的 KL 散度，使学生拟合教师的软概率分布；同时保留已知类交叉熵 CE，维持学生对真实标签的拟合。选择标准 KD 作为对照，是为了独立评估普通知识蒸馏的增量作用。",
            "3. 不确定性知识蒸馏：设教师对第 i 个样本的不确定性为 u_i，将该样本的 KL 损失乘以 exp(-u_i)，使教师较为确定的样本获得较大蒸馏权重，教师不确定样本的蒸馏影响相应减小。新增 batch-mean normalization，使加权前后的平均 KD 强度基本一致；同时分别采用 learned uncertainty head 和多次随机前向得到的 MC-epistemic uncertainty。修改目的在于降低教师不可靠预测对学生表征的影响。",
            "4. 独立 rejection 表征和 NNPU：分类分支输出已知类别概率，rejection 分支输出样本属于已知支持区域的风险分数。系统增加独立 rejection encoder 生成拒识特征，并使用 NNPU（非负正—未标记风险）进行训练：已知样本提供有标签的正类支持，mixed pool 仅作为未标记样本，不能直接将整个混合池赋予未知标签。shared-backbone 初始化复制学生 ResNet-18 的表征参数后独立训练，目的是改善 scratch CNN 拒识分支的表征学习能力。",
            "5. 不确定性拒识融合：与训练阶段的 uncertainty-KD 不同，fusion 在学生训练完成后使用。先由 known validation 拟合 uncertainty 与分类错误之间的关系，再将校准后的 uncertainty risk 与 rejection score 按 0.5/0.5 融合，用于检验 uncertainty 作为拒识信号是否比作为训练权重更合适。",
            "6. 新类候选和聚类：首先由 known validation 确定约 95% known coverage 的拒识阈值，将高风险样本形成新类候选池；随后使用 PCA 降维和 KMeans 聚类。auto K 仅依据候选特征的内部指标估计类别数，真实 K 仅用于最终评价。实验还比较 projection PCA、feature PCA 和 prototype-boundary 候选加权，以区分聚类表征因素与候选筛选因素。",
            "7. 评价方法：同时报告 AUROC、AUPR、FPR95、OSCR、known acceptance、unknown rejection、accepted-known accuracy、分数重叠、自动 K 和 ACC/NMI/ARI。unknown rejection 衡量未知样本的正确拒识比例，NMI/ARI 衡量新类聚类质量，两类指标分别对应拒识和聚类环节。"
        ],
    ),
]


PLAN = [
    "第1阶段（项目开始后第1—2个月，已完成）：完成文献和方法梳理，明确已知类、未知类、混合发现池和新类聚类的任务定义；完成 CIFAR-100 semantic-hard 60/40 数据划分和基础训练程序。",
    "第2阶段（第3—5个月，已完成）：建立 CE+standard-KD 基线，完成教师—学生蒸馏和学生特征提取；实现 uncertainty-KD，并比较 learned uncertainty 与 MC-epistemic 两种不确定性来源。",
    "第3阶段（第6—8个月，已完成）：实现独立 rejection encoder 和 NNPU 拒识，比较 scratch、random-backbone 和 shared-backbone；完成 known-validation-only 阈值校准和 mixed pool 标签流检查。",
    "第4阶段（第9—10个月，已完成主要实验）：完成 uncertainty fusion、feature PCA、自动 K 和 prototype-boundary 对照；根据 unknown rejection、known utility、分数重叠和 NMI/ARI 联合判断方法是否有效。",
    "中期检查后：固定 shared-backbone rejection + standard KD 主线，进一步研究只使用 known train/validation 信息的 hard-unknown coverage 和类条件 support 表征；不再继续扫描已被证明不稳定的 uncertainty-KD 权重和 prototype-boundary 参数。"
]


PROGRESS = [
    (
        "（一）已开展的工作",
        [
            "1. 完成研究流程的模块化设计，将实验划分为学生分类与蒸馏、拒识表征、未知候选筛选和新类聚类四个环节，以便分别评估各环节的作用。",
            "2. 完成标准 KD 基线和 uncertainty-KD 对照。标准 KD 使用 CE 维持已知类分类，使用 KL 散度蒸馏教师软概率；uncertainty-KD 进一步引入 exp(-u) 样本权重，并分别检查 learned uncertainty 和 MC-epistemic 路径的权重是否实际生效。",
            "3. 完成 rejection 表征对照。scratch CNN 在 known-only train-to-validation probe 上的准确率约为 18%—20%，shared-backbone 约为 73%—74%；据此停止使用 scratch 分支，将 shared-backbone 保留为主线配置。",
            "4. 完成 NNPU 拒识和 uncertainty fusion。NNPU 使 mixed pool 以未标记数据参与风险估计；uncertainty fusion 则在训练完成后由 known validation 校准错误风险，用于比较训练期蒸馏加权与后处理拒识融合的差异。",
            "5. 完成候选筛选与聚类对照。feature PCA 提高 novel NMI/ARI 但不改变拒识；prototype-boundary 提高排序指标却降低 calibrated unknown rejection，未能形成完整的新类发现增益。",
            "6. 完成程序检查和实验归档。当前测试为 317 passed；Python 文件 compileall 通过；各主要实验保存了数据划分、配置、checkpoint、训练日志、校准报告、discovery 报告和结果分析。"
        ],
    ),
    (
        "（二）当前困难及解决办法",
        [
            "困难一：语义相近的新类容易被误接收为某个已知类。未知样本在视觉表征上接近某个已知类别时，分类分支可能产生较高的已知类置信度；当前仍有约 74%—77% 的 unknown 分数低于 known q95，表明仅调整阈值不能消除分布重叠。下一步需改进训练得到的特征空间。",
            "困难二：uncertainty 能够排序部分已知分类错误，但未必能够区分 novel。MC-epistemic logits-KD 使 unknown rejection 平均提高 3.325 个百分点，同时伴随 NMI 和 known acceptance 下降。下一步停止扩大 uncertainty 权重搜索，将其保留为 validation-only 的拒识辅助信号。",
            "困难三：候选池规模不等同于候选池质量。prototype-boundary 选取边界附近样本后，AUROC 和 FPR95 改善，但 unknown rejection 下降 0.275 个百分点，说明该规则未能覆盖更多难识别的新类。下一步同时考察候选污染率、未知覆盖率和逐类覆盖率。",
            "困难四：聚类表征和自动类别数仍不稳定。feature PCA 能提高 NMI/ARI，但 Auto-K 误差在不同 split 上方向不一致。下一步保留 feature PCA 作为聚类模块，并将 Auto-K 误差作为独立问题研究，不以真实 K 替代自动类别数估计。"
        ],
    ),
]


STAGE_RESULTS = [
    "完成可复现实验流程。shared-backbone 在3个新划分上使AUROC平均提升0.0320、未知拒绝率提升3.61个百分点；uncertainty-KD未稳定有效，后续研究支持表征。"
]


def set_font(paragraph, size=10.5, bold=False):
    for run in paragraph.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(size)
        run.bold = bold


def clear_cell(cell):
    cell.text = ""
    return cell.paragraphs[0]


def fill_cell(cell, blocks, size=10.5):
    """Fill a cell with (heading, paragraphs) blocks, removing template guidance."""
    clear_cell(cell)
    first = True
    for heading, paragraphs in blocks:
        if not first:
            cell.add_paragraph()
        p = cell.paragraphs[0] if first else cell.paragraphs[-1]
        p.text = heading
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p, size=size, bold=True)
        first = False
        for text in paragraphs:
            p = cell.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(p, size=size)


def fill_simple_cell(cell, paragraphs, size=10.5):
    clear_cell(cell)
    for i, text in enumerate(paragraphs):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph(text)
        if i == 0:
            p.text = text
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p, size=size)


def main():
    doc = Document(TEMPLATE)

    # Cover and basic project information. Student/teacher fields remain untouched.
    doc.paragraphs[6].text = f"项目名称：{PROJECT_NAME}"
    set_font(doc.paragraphs[6], size=18, bold=True)
    doc.tables[0].cell(0, 2).text = PROJECT_NAME
    set_font(doc.tables[0].cell(0, 2).paragraphs[0], size=10.5)
    doc.tables[0].cell(1, 2).text = "□研究报告  □软件程序"
    set_font(doc.tables[0].cell(1, 2).paragraphs[0], size=10.5)

    fill_simple_cell(doc.tables[1].cell(0, 0), INTRO, size=10.5)

    research_blocks = []
    for heading, paragraphs in RESEARCH_CONTENT:
        research_blocks.append((heading, paragraphs))
    research_blocks.append(("（二）实施进度安排", PLAN))
    fill_cell(doc.tables[2].cell(0, 0), research_blocks, size=10.2)

    progress_blocks = [(heading, paragraphs) for heading, paragraphs in PROGRESS]
    fill_cell(doc.tables[3].cell(0, 0), progress_blocks, size=10.0)
    # The second row contains member division and finance guidance; leave it unchanged.

    fill_simple_cell(doc.tables[4].cell(0, 0), STAGE_RESULTS, size=10.5)

    doc.core_properties.title = PROJECT_NAME + "中期进展报告"
    doc.core_properties.subject = "实验部分已填写"
    doc.core_properties.comments = "根据项目实验记录和中期研究总结填写；学生、教师相关栏目保留模板原状。"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
