"""Generate a formal, template-compliant midterm report from the project template."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "附件2-北京交通大学大学生创新创业训练计划项目中期进展报告模板.docx"
OUTPUT = ROOT / "附件2-北京交通大学大学生创新创业训练计划项目中期进展报告_正式学术版.docx"
PROJECT_NAME = "基于不确定性知识蒸馏的新类发现方法"


INTRO = [
    "本项目面向开放集新类发现任务。在训练阶段，模型仅利用已知类别的标注样本进行分类和表征学习；在测试阶段，输入数据同时包含已知类别样本和训练阶段未纳入的未知类别样本。系统需要在缺少未知类别标注的条件下，识别未知类别样本，并对其进行聚类，从而实现新类别发现。",
    "项目以教师—学生知识蒸馏为基础，引入教师不确定性，研究其对学生表征、未知拒识和新类聚类的影响。实验采用 CIFAR-100 语义困难划分，将 60 个类别作为已知类别、40 个类别作为未知类别，并结合独立拒识表征、非负正—未标记学习和聚类模块构成流程。",
    "阶段性结果显示，教师不确定性具有一定的已知分类错误排序能力，但直接用于知识蒸馏尚未形成稳定增益。当前主要问题是语义相近的未知类别与已知类别在特征和拒识分数上的重叠，后续重点研究困难未知样本覆盖和类条件支持表征。",
]


RESEARCH_CONTENT = [
    (
        "（一）主要研究内容",
        [
            "1. 研究任务与数据协议。项目研究开放集新类发现问题。已知类别是训练阶段具有标注信息的类别；未知类别是训练阶段未纳入类别标注、但在测试阶段出现的类别；混合发现池是训练阶段仅提供图像、不提供类别标签的混合数据。实验采用 CIFAR-100 语义困难划分（semantic-hard split），每个划分包含 60 个已知类别和 40 个未知类别；混合发现池规模固定为 5,400 个样本，已知样本先验设为 0.2。",
            "2. 整体处理流程。首先由教师模型和学生模型完成已知类别分类与知识蒸馏；随后由独立拒识编码器提取拒识特征，并计算样本属于已知支持区域的风险分数；再根据已知验证集确定拒识阈值，将高风险样本形成未知类别候选池；最后在候选样本特征空间中进行聚类，并估计潜在新类别数量。模型选择依据已知验证集分类准确率，拒识阈值依据已知验证集的目标覆盖率进行校准。",
            "3. 标准知识蒸馏基线。教师模型作为参考模型，学生模型承担最终的新类发现任务。标准 logits 知识蒸馏采用温度缩放后的 KL 散度，使学生模型拟合教师模型的软类别概率分布；同时保留已知类别交叉熵损失，使学生模型继续受到真实标签监督。选择标准知识蒸馏作为基本对照，是为了单独评估知识蒸馏相对于分类监督的增量作用。",
            "4. 不确定性知识蒸馏。设教师模型对第 i 个样本的不确定性为 u_i，将该样本的 KL 蒸馏损失乘以 exp(-u_i)，使教师模型较为确定的样本获得较大的蒸馏权重，使高不确定性样本对学生参数更新的影响相应减小。实现中增加 batch-mean normalization，使加权前后的平均蒸馏强度基本一致；同时分别采用 learned uncertainty head 和多次随机前向估计的 MC-epistemic uncertainty。该修改用于检验教师不确定性是否能够降低边界样本和不可靠预测对学生表征的负面影响。",
            "5. 独立拒识表征与 NNPU。分类分支用于估计已知类别概率，独立拒识分支用于估计样本偏离已知类别支持区域的程度。系统增加独立 rejection encoder 生成拒识特征，并采用 NNPU（Non-Negative Positive-Unlabeled）风险进行训练：已知样本提供有标签的正类支持，混合发现池仅作为未标记样本参与风险估计，不将整个混合池直接赋予未知类别标签。拒识编码器采用 shared-backbone 初始化，即复制学生 ResNet-18 的表征参数后独立训练，以改善 scratch CNN 拒识分支的表征学习能力。",
            "6. 不确定性拒识融合。该方法与训练阶段的不确定性知识蒸馏不同：学生模型训练完成后，首先使用已知验证集拟合不确定性与已知分类错误之间的关系，再将校准后的不确定性风险与拒识分数按照 0.5/0.5 的固定比例融合，最后仍使用已知验证集校准拒识阈值。该实验用于区分“不确定性信号本身缺乏作用”和“不确定性用于蒸馏加权的机制不合适”两种可能性。",
            "7. 候选筛选与新类别聚类。通过已知验证集确定约 95% 已知类别覆盖率下的拒识阈值，将高风险样本作为新类别候选；候选特征经过 PCA 降维后采用 KMeans 聚类。auto K 仅依据候选特征的内部评价指标估计类别数量，真实类别数量仅用于最终回顾性评价。实验比较 projection PCA、feature PCA 和 prototype-boundary 候选加权，分别分析聚类表征和候选筛选对新类发现的影响。",
            "8. 评价指标。拒识部分同时报告 AUROC、AUPR、FPR95、OSCR、已知类别接受率、未知类别拒识率、接受后已知类别准确率和分数重叠；聚类部分报告自动 K 误差以及 ACC、NMI、ARI。未知类别拒识率衡量未知样本被正确拒绝的比例，NMI 和 ARI 衡量候选样本聚类结果与真实新类别划分的一致程度，分别对应拒识环节和聚类环节。",
        ],
    ),
]


PLAN = [
    "第 1 阶段（项目开始后第 1—2 个月，已完成）：完成相关研究方法梳理，明确开放集新类发现任务、已知类别、未知类别和混合发现池的定义；完成 CIFAR-100 语义困难 60/40 数据划分和基础训练程序。",
    "第 2 阶段（第 3—5 个月，已完成）：建立交叉熵与标准知识蒸馏基线，完成教师—学生模型训练和学生特征提取；实现不确定性知识蒸馏，并比较 learned uncertainty 与 MC-epistemic 两种不确定性来源。",
    "第 3 阶段（第 6—8 个月，已完成）：实现独立拒识编码器和 NNPU 拒识目标，比较 scratch、random-backbone 和 shared-backbone 初始化方式；完成已知验证集阈值校准和混合发现池标签流检查。",
    "第 4 阶段（第 9—10 个月，已完成主要实验）：完成 validation-only 不确定性拒识融合、feature PCA、自动 K 和 prototype-boundary 候选加权对照；依据拒识、已知类别效用、分数重叠和聚类指标进行联合评价。",
    "中期检查后：固定 shared-backbone 拒识表征与标准知识蒸馏作为主线配置，研究仅依赖已知训练集、已知验证集和模型预测信息的困难未知样本覆盖与类条件支持表征；停止对已显示不稳定的 uncertainty-KD 权重和 prototype-boundary 参数进行继续搜索。",
]


PROGRESS = [
    (
        "（一）已开展的工作",
        [
            "1. 完成实验流程和程序结构审阅，将研究流程划分为学生分类与蒸馏、拒识表征、未知候选筛选和新类别聚类四个环节，并为各环节分别设置对照指标。",
            "2. 完成标准知识蒸馏基线和不确定性知识蒸馏对照。标准方法使用已知类别交叉熵和教师 logits 蒸馏；不确定性方法增加 exp(-u) 样本权重，并检查 learned uncertainty 和 MC-epistemic 两条路径的权重、损失及梯度是否实际生效。",
            "3. 完成拒识表征初始化对照。scratch CNN 在已知类别训练—验证探针任务上的准确率约为 18%—20%，shared-backbone 约为 73%—74%；据此停止使用 scratch 拒识分支，将 shared-backbone 作为当前主线配置。2081—2083 三个新划分中，shared-backbone 相对 scratch 的平均 AUROC 提升 0.0320，OSCR 提升 0.0280，未知类别拒识率提升 3.61 个百分点。",
            "4. 完成 NNPU 拒识和 validation-only 不确定性融合实验。NNPU 使混合发现池以未标记数据参与风险估计；不确定性融合在学生模型训练完成后由已知验证集校准错误风险，用于比较训练阶段蒸馏加权与后处理拒识融合的差异。7 个新划分中，融合方法平均使 AUROC 提升 0.011677，未知类别拒识率提升 4.393 个百分点，但不同划分之间仍存在波动。",
            "5. 完成候选筛选和聚类对照。feature PCA 在 2110、2111、2117 三个新划分中使全体 novel NMI 平均提升 0.06896、ARI 平均提升 0.07846，但不改变拒识指标；prototype-boundary 在 2119 中使 AUROC 从 0.778810 提升至 0.789709、FPR95 从 64.717% 降至 58.750%，但未知类别拒识率从 25.325% 降至 25.050%，未达到预设联合标准。",
            "6. 完成程序与实验归档检查。当前单元测试为 317 passed；Python 文件 compileall 通过；主要实验均保存数据划分、配置文件、模型 checkpoint、训练日志、校准报告、discovery 报告和结果分析文件。",
        ],
    ),
    (
        "（二）当前困难及解决办法",
        [
            "困难一：语义相近的未知类别与已知类别在特征空间和拒识分数上存在明显重叠。当前仍有约 74%—77% 的未知样本分数低于已知类别 q95，说明仅调整拒识阈值不能从根本上消除重叠。解决思路是改进训练阶段的特征空间和类条件支持区域，而不是继续进行单纯阈值搜索。",
            "困难二：不确定性可以排序部分已知分类错误，但尚未稳定区分未知类别。MC-epistemic logits-KD 在 2114—2116 三个新划分上使未知类别拒识率平均提高 3.325 个百分点，但整体 NMI 平均下降 0.008623，已知类别接受率平均下降 0.872 个百分点。解决办法是停止不确定性权重函数搜索，将不确定性保留为 validation-only 拒识辅助信号。",
            "困难三：候选池规模扩大或改变候选规则并不必然提高困难未知样本覆盖率。prototype-boundary 在 2119 中改善排序指标，但未知类别拒识率下降 0.275 个百分点，表明候选规则未形成完整任务收益。后续将同时评价候选污染率、未知覆盖率和逐类别覆盖率，避免仅依据排序指标调整候选规则。",
            "困难四：自动类别数估计和聚类表示仍具有划分敏感性。feature PCA 能够改善 NMI 和 ARI，但 Auto-K 误差在不同划分上方向不一致。后续将 feature PCA 作为聚类模块保留，并将 Auto-K 误差作为独立问题研究，不以真实类别数量替代自动估计结果。",
        ],
    ),
]


STAGE_RESULTS = [
    "完成可复现实验流程。shared-backbone 在3个新划分中使AUROC平均提升0.0320、未知类别拒识率提升3.61个百分点；不确定性知识蒸馏未形成稳定增益。"
]


def set_font(paragraph, size=10.5, bold=False):
    for run in paragraph.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(size)
        run.bold = bold


def clear_cell(cell):
    cell.text = ""
    return cell.paragraphs[0]


def fill_cell(cell, blocks, size=10.5):
    clear_cell(cell)
    first = True
    for heading, paragraphs in blocks:
        if not first:
            cell.add_paragraph()
        p = cell.paragraphs[0] if first else cell.paragraphs[-1]
        p.text = heading
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p, size=size, bold=True)
        first = False
        for text in paragraphs:
            p = cell.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(p, size=size)


def fill_simple_cell(cell, paragraphs, size=10.5):
    clear_cell(cell)
    for i, text in enumerate(paragraphs):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph(text)
        if i == 0:
            p.text = text
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p, size=size)


def main():
    doc = Document(TEMPLATE)
    doc.paragraphs[6].text = f"项目名称：{PROJECT_NAME}"
    set_font(doc.paragraphs[6], size=18, bold=True)
    doc.tables[0].cell(0, 2).text = PROJECT_NAME
    set_font(doc.tables[0].cell(0, 2).paragraphs[0], size=10.5)
    doc.tables[0].cell(1, 2).text = "□研究报告  □软件程序"
    set_font(doc.tables[0].cell(1, 2).paragraphs[0], size=10.5)

    fill_simple_cell(doc.tables[1].cell(0, 0), INTRO, size=10.5)
    fill_cell(doc.tables[2].cell(0, 0), RESEARCH_CONTENT + [("（二）实施进度安排", PLAN)], size=10.2)
    fill_cell(doc.tables[3].cell(0, 0), PROGRESS, size=10.0)
    fill_simple_cell(doc.tables[4].cell(0, 0), STAGE_RESULTS, size=10.5)

    doc.core_properties.title = PROJECT_NAME + "中期进展报告"
    doc.core_properties.subject = "实验部分正式学术表述版"
    doc.core_properties.comments = "依据项目实验记录填写；学生、成员和指导教师相关栏目保持模板原状。"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
